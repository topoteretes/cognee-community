"""Content extraction: text decoding, PDF and Word parsing, and skipping bad files."""

import io

import pytest
from _fake_egnyte import FakeEgnyte

from cognee_community_connector_egnyte import egnyte as eg


def _pdf_with_text(text):
    """Build a minimal one-page PDF whose content stream draws ``text``."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
    out.write(f"{trailer}startxref\n{xref_at}\n%%EOF\n".encode())
    return out.getvalue()


def _docx_with(paragraph, table_row):
    from docx import Document

    document = Document()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=len(table_row))
    for cell, value in zip(table.rows[0].cells, table_row, strict=True):
        cell.text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _extract(path, data):
    fake = FakeEgnyte()
    fake.put(path, data)
    return eg.extract_file_content(fake, path)


def test_text_files_are_decoded_as_utf8():
    assert _extract("/Shared/notes.md", "Café meeting notes".encode()) == "Café meeting notes"


def test_pdf_text_is_extracted():
    pdf = _pdf_with_text("Quarterly revenue grew")
    assert "Quarterly revenue grew" in _extract("/Shared/report.pdf", pdf)


def test_docx_paragraphs_and_tables_are_extracted():
    docx = _docx_with("Payment within 30 days.", ["Fee", "$5,000"])
    text = _extract("/Shared/contract.docx", docx)

    assert "Payment within 30 days." in text
    assert "Fee | $5,000" in text


@pytest.mark.parametrize("path", ["/Shared/broken.pdf", "/Shared/broken.docx"])
def test_unparseable_files_are_skipped(path):
    assert _extract(path, b"this is not a valid document") is None


@pytest.mark.parametrize(
    ("path", "supported"),
    [
        ("/Shared/a.TXT", True),
        ("/Shared/a.pdf", True),
        ("/Shared/a.docx", True),
        ("/Shared/a.csv", True),
        ("/Shared/a.png", False),
        ("/Shared/a.doc", False),
        ("/Shared/README", False),
    ],
)
def test_supported_extensions(path, supported):
    assert eg.is_supported_file(path) is supported
