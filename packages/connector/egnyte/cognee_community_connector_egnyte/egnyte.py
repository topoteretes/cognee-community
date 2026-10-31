"""Egnyte connector for cognee: a ``dlt`` source that turns Egnyte folders into memory.

Sync one or more Egnyte folders (plain text, Markdown, CSV, JSON, PDF and Word
files) into cognee, incrementally and with forget-on-deletion.  The resource
produced here is handed directly to :func:`cognee.remember`::

    import cognee
    from cognee_community_connector_egnyte import egnyte_source

    await cognee.remember(
        egnyte_source(domain="acme", folder_paths=["/Shared/Projects"]),
        dataset_name="egnyte_projects",
        primary_key="id",
        write_disposition="merge",   # incremental upsert by file group id
        max_rows_per_table=0,        # 0 = no row cap (folders often exceed the default 50)
    )

Design
------
* **Auth**: an OAuth access token for the Egnyte Public API, sent as a Bearer
  header.  Get it once and pass it as ``access_token``: with the authorization
  code flow for "public application" keys (:func:`authorization_url`,
  :func:`code_from_redirect`, :func:`exchange_code`, see
  ``examples/login.py``), or with the password grant for "internal
  application" keys (:func:`get_access_token`, which the source also runs
  itself when given an API key, username and password).  The connector only
  issues ``GET`` requests.
* **Primary key**: the file's ``group_id``.  Egnyte keeps the same group id
  across versions, renames and moves (a new version only changes the
  ``entry_id``), so a renamed file is updated in place instead of looking like a
  delete plus a new file.
* **Incremental cursor**: the Events API id.  The first run reads the latest
  event id, then walks the configured folders.  Later runs read the file system
  events since that id and re-check only the paths those events touched.
* **Forget-on-delete**: the resource state keeps a ``group_id -> path`` map of
  every file in scope.  A touched path that no longer resolves to a live file
  (deleted, or moved out of scope) is emitted with the ``_deleted`` hard-delete
  marker; dlt drops it on ``merge`` and cognee's ``orphan_cleanup`` purges it
  from the graph, vector and relational stores.  Deleting or moving a folder
  emits a single folder event, so the map is also how the files under it are
  found.
* **Stale cursor**: Egnyte only keeps a limited window of events.  If the saved
  cursor is older than the oldest event still available, the connector falls
  back to a full walk and tombstones every previously known file that is gone,
  instead of silently missing changes.
* **Rate limits**: default developer keys allow 2 calls per second and 1,000
  calls per day.  Requests are paced to ``requests_per_second`` and the
  per-second throttle is retried; hitting the daily quota raises an error, so
  the run fails without advancing the cursor and the next run picks up where
  this one left off.
* **Isolation**: each resource name gets its own dlt pipeline state and its own
  cleanup scope, so two Egnyte syncs (different domains or folder sets) can
  share a dataset without overwriting each other's cursor or forgetting each
  other's files, as long as each uses its own ``resource_name``.

Limitations
-----------
* Only files with a supported extension are ingested (see
  ``SUPPORTED_EXTENSIONS``).  Other files are skipped with a warning.
* A file whose content cannot be parsed is skipped with a warning rather than
  failing the whole sync, so an earlier ingested version stays in memory until
  the file changes again or is deleted.
"""

import io
import os
import time
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from typing import Any
from urllib.parse import parse_qs, quote, urlencode, urlparse

from cognee.shared.logging_utils import get_logger
from cognee.tasks.ingestion import dlt_utils
from cognee.tasks.ingestion.dlt_utils import DOCUMENT_SOURCE_ATTR

logger = get_logger("egnyte_connector")

EGNYTE_SOURCE_NAME = "egnyte"
EGNYTE_TABLE_NAME = "egnyte_files"

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".csv",
    ".tsv",
    ".json",
    ".yaml",
    ".yml",
    ".log",
}
PDF_EXTENSION = ".pdf"
DOCX_EXTENSION = ".docx"
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {PDF_EXTENSION, DOCX_EXTENSION}

# Egnyte returns at most this many entries per folder listing / events page.
_PAGE_SIZE = 100
# Retry budget for throttled (per-second quota) and transient (5xx / network) failures.
_MAX_RETRIES = 5
_REQUEST_TIMEOUT_SECONDS = 60

# Egnyte reports quota errors as HTTP 403 plus this header, not as HTTP 429.
_ERROR_CODE_HEADER = "X-Mashery-Error-Code"
_OVER_QPS = "ERR_403_DEVELOPER_OVER_QPS"
_OVER_RATE = "ERR_403_DEVELOPER_OVER_RATE"


class EgnyteError(RuntimeError):
    """An Egnyte API call failed."""


class EgnyteNotFoundError(EgnyteError):
    """The requested file or folder does not exist (HTTP 404)."""


class EgnyteQuotaExceededError(EgnyteError):
    """The API key's daily call quota is used up."""


@dataclass(frozen=True)
class _EgnyteConfig:
    domain: str
    folder_paths: tuple[str, ...]
    max_file_size_mb: int


# ---------------------------------------------------------------------------
# HTTP client
# ---------------------------------------------------------------------------
def _host(domain: str) -> str:
    """Accept either ``acme`` or ``acme.egnyte.com``."""
    domain = domain.strip().removeprefix("https://").rstrip("/")
    return domain if "." in domain else f"{domain}.egnyte.com"


def _encode_path(path: str) -> str:
    return quote(path, safe="/")


DEFAULT_OAUTH_SCOPE = "Egnyte.filesystem"
DEFAULT_REDIRECT_URI = "https://localhost/callback"


def authorization_url(
    domain: str,
    client_id: str,
    state: str,
    *,
    redirect_uri: str = DEFAULT_REDIRECT_URI,
    scope: str = DEFAULT_OAUTH_SCOPE,
) -> str:
    """Return the page where an Egnyte user approves access for a public app.

    This is step one of the OAuth authorization code flow.  Egnyte then
    redirects the browser to ``redirect_uri`` with a ``code`` and the same
    ``state``; pass that URL to :func:`code_from_redirect` and the code to
    :func:`exchange_code`.  ``redirect_uri`` must match the one registered with
    the API key.
    """
    query = urlencode(
        {
            "client_id": client_id,
            "redirect_uri": redirect_uri,
            "scope": scope,
            "state": state,
            "response_type": "code",
        }
    )
    return f"https://{_host(domain)}/puboauth/token?{query}"


def code_from_redirect(redirect_url: str, state: str) -> str:
    """Pull the authorization code out of the URL Egnyte redirected to.

    Checks ``state`` so a code from a sign-in this process did not start is
    never accepted.
    """
    params = parse_qs(urlparse(redirect_url.strip()).query)
    if "error" in params:
        detail = params.get("error_description", params["error"])[0]
        raise EgnyteError(f"Egnyte: sign-in was not approved: {detail}")
    if params.get("state", [None])[0] != state:
        raise EgnyteError("Egnyte: the redirect URL is from a different sign-in (state mismatch).")
    code = params.get("code", [None])[0]
    if not code:
        raise EgnyteError("Egnyte: the redirect URL has no authorization code.")
    return code


def exchange_code(
    domain: str,
    client_id: str,
    client_secret: str,
    code: str,
    *,
    redirect_uri: str = DEFAULT_REDIRECT_URI,
    session: Any = None,
) -> str:
    """Trade an authorization code for an access token (public apps)."""
    return _request_token(
        domain,
        {
            "client_id": client_id,
            "client_secret": client_secret,
            "redirect_uri": redirect_uri,
            "code": code,
            "grant_type": "authorization_code",
        },
        session,
    )


def get_access_token(
    domain: str,
    api_key: str,
    username: str,
    password: str,
    session: Any = None,
) -> str:
    """Request an access token with the OAuth password grant.

    This is the flow for Egnyte "internal application" API keys.  Egnyte limits
    token requests per user per hour, so request a token once, store it (for
    example in ``EGNYTE_ACCESS_TOKEN``) and reuse it.
    """
    return _request_token(
        domain,
        {
            "client_id": api_key,
            "username": username,
            "password": password,
            "grant_type": "password",
        },
        session,
    )


def _request_token(domain: str, data: dict, session: Any = None) -> str:
    if session is None:
        import requests

        session = requests.Session()

    response = session.post(
        f"https://{_host(domain)}/puboauth/token",
        data=data,
        timeout=_REQUEST_TIMEOUT_SECONDS,
    )
    if response.status_code != 200:
        raise EgnyteError(
            f"Egnyte: token request failed with HTTP {response.status_code}: {response.text}"
        )
    return response.json()["access_token"]


class EgnyteClient:
    """Minimal read-only client for the Egnyte Public API.

    ``session`` only needs a ``requests``-style ``get(url, params=..., headers=...,
    timeout=...)``, which keeps the client easy to fake in tests.  ``sleep`` and
    ``clock`` are injectable for the same reason.
    """

    def __init__(
        self,
        domain: str,
        access_token: str,
        *,
        session: Any = None,
        requests_per_second: float = 2.0,
        sleep: Callable[[float], None] = time.sleep,
        clock: Callable[[], float] = time.monotonic,
    ):
        if session is None:
            import requests

            session = requests.Session()

        self._base_url = f"https://{_host(domain)}/pubapi"
        self._headers = {"Authorization": f"Bearer {access_token}"}
        self._session = session
        self._min_interval = 1.0 / requests_per_second if requests_per_second > 0 else 0.0
        self._sleep = sleep
        self._clock = clock
        self._last_request_at: float | None = None

    # -- public API ---------------------------------------------------------
    def get_event_cursor(self) -> dict[str, Any]:
        """Return ``{"latest_event_id": ..., "oldest_event_id": ...}``."""
        return self._get_json("/v1/events/cursor")

    def list_events(self, start_id: int, folder: str) -> list[dict[str, Any]]:
        """Return file system events after ``start_id`` under ``folder``.

        Egnyte answers 204 No Content when there are no newer events.
        """
        response = self._request(
            "/v1/events",
            params={
                "id": start_id,
                "folder": folder,
                "type": "file_system",
                "count": _PAGE_SIZE,
            },
        )
        if response.status_code == 204:
            return []
        return response.json().get("events", [])

    def list_folder(self, path: str, offset: int = 0) -> dict[str, Any]:
        """Return one page of a folder listing (``files`` and ``folders``)."""
        return self._get_json(
            f"/v1/fs{_encode_path(path)}",
            params={"list_content": "true", "count": _PAGE_SIZE, "offset": offset},
        )

    def get_metadata(self, path: str) -> dict[str, Any]:
        """Return metadata for a file or folder; raises EgnyteNotFoundError if gone."""
        return self._get_json(f"/v1/fs{_encode_path(path)}")

    def download(self, path: str) -> bytes:
        """Download the current version of a file."""
        return self._request(f"/v1/fs-content{_encode_path(path)}").content

    # -- internals ----------------------------------------------------------
    def _get_json(self, endpoint: str, params: dict | None = None) -> dict[str, Any]:
        return self._request(endpoint, params=params).json()

    def _request(self, endpoint: str, params: dict | None = None):
        url = f"{self._base_url}{endpoint}"
        for attempt in range(_MAX_RETRIES + 1):
            self._wait_for_rate_limit()
            try:
                response = self._session.get(
                    url, params=params, headers=self._headers, timeout=_REQUEST_TIMEOUT_SECONDS
                )
            except OSError as exc:
                # requests' ConnectionError and Timeout are OSError subclasses.
                if attempt == _MAX_RETRIES:
                    raise EgnyteError(f"Egnyte: request to {endpoint} failed: {exc}") from exc
                self._sleep(2**attempt)
                continue

            status = response.status_code
            error_code = response.headers.get(_ERROR_CODE_HEADER, "")

            if status == 403 and error_code == _OVER_QPS and attempt < _MAX_RETRIES:
                self._sleep(_retry_after(response.headers, attempt))
                continue
            if status >= 500 and attempt < _MAX_RETRIES:
                self._sleep(2**attempt)
                continue

            if status in (200, 204):
                return response
            if status == 404:
                raise EgnyteNotFoundError(f"Egnyte: {endpoint} not found")
            if status == 403 and error_code == _OVER_RATE:
                raise EgnyteQuotaExceededError(
                    "Egnyte: the API key's daily call quota is used up. The sync stopped "
                    "without advancing its cursor; re-run it after the quota resets."
                )
            if status == 401:
                raise EgnyteError(
                    "Egnyte: the access token was rejected (HTTP 401). Check EGNYTE_ACCESS_TOKEN "
                    "or the API key, username and password."
                )
            raise EgnyteError(
                f"Egnyte: GET {endpoint} failed with HTTP {status}: {response.text[:500]}"
            )

        raise EgnyteError(f"Egnyte: GET {endpoint} still throttled after {_MAX_RETRIES} retries")

    def _wait_for_rate_limit(self) -> None:
        if self._min_interval and self._last_request_at is not None:
            elapsed = self._clock() - self._last_request_at
            if elapsed < self._min_interval:
                self._sleep(self._min_interval - elapsed)
        self._last_request_at = self._clock()


def _retry_after(headers, attempt: int) -> float:
    try:
        return max(float(headers.get("Retry-After", "")), 0.0)
    except ValueError:
        return float(2**attempt)


# ---------------------------------------------------------------------------
# Content extraction (dispatch on file extension)
# ---------------------------------------------------------------------------
def _extension(path: str) -> str:
    return os.path.splitext(path)[1].lower()


def is_supported_file(path: str) -> bool:
    return _extension(path) in SUPPORTED_EXTENSIONS


def extract_file_content(client: EgnyteClient, path: str) -> str | None:
    """Return extracted text for an Egnyte file, or None to skip it.

    A file that can't be parsed (corrupt PDF, unreadable Word file) is logged and
    skipped rather than aborting the whole sync.  Auth, quota and connectivity
    failures are raised, because they affect every file.
    """
    extension = _extension(path)
    try:
        data = client.download(path)
    except EgnyteNotFoundError:
        # Deleted between the listing and the download; its delete event will
        # be handled by the next incremental run.
        logger.info("Skipping Egnyte file '%s': it was deleted during the sync.", path)
        return None
    try:
        if extension == PDF_EXTENSION:
            return _extract_pdf_text(data)
        if extension == DOCX_EXTENSION:
            return _extract_docx_text(data)
        return data.decode("utf-8", errors="replace")
    except Exception as exc:
        logger.warning("Skipping Egnyte file '%s': content extraction failed: %s", path, exc)
        return None


def _extract_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=False)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Scope helpers (Egnyte paths are case-insensitive)
# ---------------------------------------------------------------------------
def _normalize_folder(path: str) -> str:
    path = "/" + path.strip().strip("/")
    return path


def _is_under(path: str, folder: str) -> bool:
    path_key = path.casefold()
    folder_key = folder.rstrip("/").casefold()
    return path_key == folder_key or path_key.startswith(folder_key + "/")


def _in_scope(path: str, config: _EgnyteConfig) -> bool:
    return any(_is_under(path, root) for root in config.folder_paths)


# ---------------------------------------------------------------------------
# Sync state machine (pure given a client + state dict, so it is unit-testable)
# ---------------------------------------------------------------------------
def _iter_rows(client: EgnyteClient, config: _EgnyteConfig, state: dict) -> Iterator[dict]:
    """Yield one row per new or changed file in scope, plus hard-delete tombstones.

    ``state`` stands in for dlt's resource state: ``{"cursor": int, "files":
    {group_id: path}}``.  It is only updated after every row has been yielded,
    so a failed run leaves it untouched and the next run retries the same window.
    """
    cursor_info = client.get_event_cursor()
    latest_event_id = int(cursor_info["latest_event_id"])
    oldest_event_id = int(cursor_info.get("oldest_event_id", 0))
    known_files: dict[str, str] = dict(state.get("files") or {})
    saved_cursor = state.get("cursor")

    if saved_cursor is None:
        # Read the cursor BEFORE walking, so a file changed during a slow first
        # sync is picked up by the next incremental run rather than missed.
        files = yield from _full_sync(client, config, known_files)
        logger.info("Egnyte: initial sync found %d file(s).", len(files))
    elif int(saved_cursor) < oldest_event_id:
        logger.warning(
            "Egnyte: saved event cursor %s is older than the oldest available event %s; "
            "running a full sync to catch up.",
            saved_cursor,
            oldest_event_id,
        )
        files = yield from _full_sync(client, config, known_files)
    else:
        files = yield from _incremental_sync(
            client, config, known_files, int(saved_cursor), latest_event_id
        )

    state["files"] = files
    state["cursor"] = latest_event_id


def _full_sync(client, config: _EgnyteConfig, known_files: dict[str, str]):
    """Walk every configured folder; tombstone known files that no longer exist."""
    files: dict[str, str] = {}
    for root in config.folder_paths:
        try:
            for file_meta in _walk_folder(client, root):
                row = _file_to_row(client, file_meta, config)
                files[file_meta["group_id"]] = file_meta["path"]
                if row is not None:
                    yield row
        except EgnyteNotFoundError as exc:
            # Fail instead of treating a typo'd or unshared folder as empty,
            # which would forget everything previously synced from it.
            raise EgnyteError(
                f"Egnyte: folder '{root}' was not found, or this user can't see it."
            ) from exc

    for group_id in known_files.keys() - files.keys():
        yield {"id": group_id, "_deleted": True}
    return files


def _incremental_sync(
    client,
    config: _EgnyteConfig,
    known_files: dict[str, str],
    cursor: int,
    latest_event_id: int,
):
    """Re-check only the paths touched by events in ``(cursor, latest_event_id]``."""
    touched_files: set[str] = set()
    touched_folders: set[str] = set()
    for event in _list_events(client, config, cursor, latest_event_id):
        data = event.get("data") or {}
        paths = [p for p in (data.get("source_path"), data.get("target_path")) if p]
        if data.get("is_folder"):
            touched_folders.update(paths)
        else:
            touched_files.update(paths)

    files = dict(known_files)
    live: dict[str, dict] = {}
    gone_folders: list[str] = []

    # A folder event stands for every file under it: files we knew under the
    # old path must be re-checked, and the folder at its new path is walked.
    for folder in touched_folders:
        touched_files.update(path for path in known_files.values() if _is_under(path, folder))
        if not _in_scope(folder, config):
            continue
        try:
            for file_meta in _walk_folder(client, folder):
                live[file_meta["group_id"]] = file_meta
        except EgnyteNotFoundError:
            # Deleted, or moved away. Its files can't exist at their old paths,
            # so don't spend one metadata call per file finding that out.
            gone_folders.append(folder)

    live_paths = {meta["path"].casefold() for meta in live.values()}
    for path in touched_files:
        if not _in_scope(path, config) or path.casefold() in live_paths:
            continue
        if any(_is_under(path, folder) for folder in gone_folders):
            continue
        try:
            file_meta = client.get_metadata(path)
        except EgnyteNotFoundError:
            continue
        if not file_meta.get("is_folder"):
            live[file_meta["group_id"]] = file_meta

    touched_keys = {path.casefold() for path in touched_files}
    for group_id, path in known_files.items():
        if path.casefold() in touched_keys and group_id not in live:
            del files[group_id]
            yield {"id": group_id, "_deleted": True}

    for group_id, file_meta in live.items():
        files[group_id] = file_meta["path"]
        row = _file_to_row(client, file_meta, config)
        if row is not None:
            yield row

    logger.info(
        "Egnyte: incremental sync re-checked %d path(s) and %d folder(s).",
        len(touched_files),
        len(touched_folders),
    )
    return files


def _list_events(client, config: _EgnyteConfig, cursor: int, latest_event_id: int):
    """Yield events after ``cursor`` up to ``latest_event_id`` for every root folder.

    Events newer than ``latest_event_id`` are left for the next run, so polling
    several folders one after another can't skip an event that lands mid-run.
    """
    seen: set[int] = set()
    for root in config.folder_paths:
        start_id = cursor
        while start_id < latest_event_id:
            page = client.list_events(start_id, folder=root)
            if not page:
                break
            for event in page:
                event_id = int(event["id"])
                if event_id <= latest_event_id and event_id not in seen:
                    seen.add(event_id)
                    yield event
            start_id = max(int(event["id"]) for event in page)
            if len(page) < _PAGE_SIZE:
                break


def _walk_folder(client, root: str) -> Iterator[dict]:
    """Yield metadata for every file under ``root``, recursively (breadth first).

    Raises EgnyteNotFoundError if ``root`` itself is missing.  A subfolder that
    disappears mid-walk is skipped; its delete event is handled on the next run.
    """
    queue = [root]
    seen_folders: set[str] = set()
    while queue:
        folder = queue.pop(0)
        if folder.casefold() in seen_folders:
            continue
        seen_folders.add(folder.casefold())

        offset = 0
        while True:
            try:
                listing = client.list_folder(folder, offset=offset)
            except EgnyteNotFoundError:
                if folder is root:
                    raise
                break
            files = listing.get("files") or []
            folders = listing.get("folders") or []
            yield from files
            queue.extend(sub["path"] for sub in folders)

            returned = len(files) + len(folders)
            if returned < _PAGE_SIZE:
                break
            offset += returned


def _file_to_row(client, file_meta: dict[str, Any], config: _EgnyteConfig) -> dict | None:
    path = file_meta["path"]
    group_id = file_meta["group_id"]

    if not is_supported_file(path):
        logger.warning("Skipping unsupported Egnyte file '%s'.", path)
        return None

    size = file_meta.get("size")
    if size and int(size) > config.max_file_size_mb * 1024 * 1024:
        logger.warning(
            "Skipping Egnyte file '%s': size exceeds max_file_size_mb=%d.",
            path,
            config.max_file_size_mb,
        )
        return None

    content = extract_file_content(client, path)
    if content is None or not content.strip():
        return None

    # Document-mode row contract: {id, title, content, url}.  The title is the
    # full path so the folder structure ("/Shared/Clients/Acme/...") is part of
    # what cognify sees.
    return {
        "id": group_id,
        "title": path,
        "content": content,
        "url": f"https://{_host(config.domain)}/navigate/file/{group_id}",
        "_deleted": False,
    }


# ---------------------------------------------------------------------------
# Public factory
# ---------------------------------------------------------------------------
def _resolve_folder_paths(folder_paths: list[str] | str | None) -> tuple[str, ...]:
    if folder_paths is None:
        folder_paths = os.getenv("EGNYTE_FOLDER_PATHS", "")
    if isinstance(folder_paths, str):
        folder_paths = folder_paths.split(",")
    resolved = tuple(_normalize_folder(p) for p in folder_paths if p.strip())
    if not resolved:
        raise ValueError(
            "folder_paths is required (pass a list such as ['/Shared/Projects'] or set "
            "EGNYTE_FOLDER_PATHS to a comma-separated list)."
        )
    return resolved


def _resolve_credentials(
    access_token: str | None,
    api_key: str | None,
    username: str | None,
    password: str | None,
) -> tuple[str | None, tuple[str, str, str] | None]:
    """Fill credentials from the environment: a token, or the password grant inputs.

    Raises ``ValueError`` when neither auth method is complete.  Makes no request.
    """
    token = access_token or os.getenv("EGNYTE_ACCESS_TOKEN")
    if token:
        return token, None
    grant = (
        api_key or os.getenv("EGNYTE_API_KEY"),
        username or os.getenv("EGNYTE_USERNAME"),
        password or os.getenv("EGNYTE_PASSWORD"),
    )
    if not all(grant):
        raise ValueError(
            "Egnyte credentials required: set EGNYTE_ACCESS_TOKEN, or EGNYTE_API_KEY "
            "plus EGNYTE_USERNAME and EGNYTE_PASSWORD."
        )
    return None, grant


def _build_client(
    domain: str,
    access_token: str | None,
    api_key: str | None,
    username: str | None,
    password: str | None,
    requests_per_second: float,
) -> EgnyteClient:
    token, grant = _resolve_credentials(access_token, api_key, username, password)
    if token is None:
        token = get_access_token(domain, *grant)
    return EgnyteClient(domain, token, requests_per_second=requests_per_second)


def egnyte_source(
    domain: str | None = None,
    folder_paths: list[str] | str | None = None,
    *,
    access_token: str | None = None,
    api_key: str | None = None,
    username: str | None = None,
    password: str | None = None,
    max_file_size_mb: int | None = None,
    requests_per_second: float | None = None,
    resource_name: str = EGNYTE_TABLE_NAME,
    check_active: Callable[[], None] | None = None,
    client: EgnyteClient | None = None,
):
    """Return a ``dlt`` resource yielding one row per in-scope Egnyte file.

    Any argument left as ``None`` falls back to the matching ``EGNYTE_*``
    environment variable.  Hand the result to ``cognee.remember(...)`` with
    ``write_disposition="merge"`` and ``primary_key="id"``.

    Args:
        domain: Egnyte domain, either ``acme`` or ``acme.egnyte.com``.
        folder_paths: Folders to sync recursively, e.g. ``["/Shared/Projects"]``.
        access_token: OAuth access token.  Preferred over the password grant.
        api_key: API key of an "internal application", used with ``username``
            and ``password`` to request a token when no ``access_token`` is set.
        username: Egnyte username for the password grant.
        password: Egnyte password for the password grant.
        max_file_size_mb: Skip files larger than this (default 25).
        requests_per_second: Request pacing (default 2, Egnyte's default key limit).
        resource_name: Stable dlt resource name.  Give each domain or folder set
            its own name when several sync into one dataset, so their cursors
            and deletes stay separate.
        check_active: Optional host authorization checkpoint during extraction.
        client: Pre-built :class:`EgnyteClient`.  Mainly an injection point for
            tests; when omitted one is built from the settings above.
    """
    try:
        import dlt
    except ImportError as exc:
        raise ImportError(
            'The Egnyte connector requires dlt: pip install "cognee-community-connector-egnyte".'
        ) from exc

    if getattr(dlt_utils, "DOCUMENT_SYNC_VERSION", 0) < 1:
        raise RuntimeError(
            "Egnyte sync requires a Cognee build with table-scoped DLT document "
            "cleanup. Upgrade Cognee before syncing to avoid deleting another source's data."
        )
    if not resource_name:
        raise ValueError("resource_name must be a non-empty string.")

    resolved_domain = domain or os.getenv("EGNYTE_DOMAIN")
    if not resolved_domain:
        raise ValueError("domain is required (pass it explicitly or set EGNYTE_DOMAIN).")
    if client is None:
        # Fail when the source is built, not halfway through a pipeline run.
        _resolve_credentials(access_token, api_key, username, password)

    config = _EgnyteConfig(
        domain=resolved_domain,
        folder_paths=_resolve_folder_paths(folder_paths),
        max_file_size_mb=(
            max_file_size_mb
            if max_file_size_mb is not None
            else int(os.getenv("EGNYTE_MAX_FILE_SIZE_MB", "25"))
        ),
    )
    rate = (
        requests_per_second
        if requests_per_second is not None
        else float(os.getenv("EGNYTE_REQUESTS_PER_SECOND", "2"))
    )

    @dlt.resource(
        name=resource_name,
        write_disposition="merge",
        primary_key="id",
        # `_deleted` is a boolean hard-delete marker (matching the google-drive
        # connector): rows where it is True are removed from the dlt destination
        # on merge, which propagates the deletion through cognee's orphan_cleanup.
        columns={"_deleted": {"data_type": "bool", "hard_delete": True}},
    )
    def egnyte_files():
        egnyte = client or _build_client(
            config.domain, access_token, api_key, username, password, rate
        )
        yield from dlt_utils.guarded_rows(
            _iter_rows(egnyte, config, dlt.current.resource_state()),
            check_active,
        )

    resource = egnyte_files()
    # Opt into the document ingestion path: each file row (id/title/content/url)
    # becomes a text document that flows through normal cognify (LLM graph
    # extraction) instead of the relational dlt-row path.
    setattr(resource, DOCUMENT_SOURCE_ATTR, EGNYTE_SOURCE_NAME)
    # Give this resource its own pipeline state and cleanup scope, so syncs
    # into different datasets, or other Egnyte syncs in the same dataset,
    # never share a cursor or forget each other's files.
    setattr(resource, dlt_utils.PIPELINE_SCOPE_ATTR, resource_name)
    return resource
